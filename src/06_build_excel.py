# -*- coding: utf-8 -*-
"""生成 docs/飞鹤投诉分拣案例_数据与测算.xlsx：说明 / 数据 / 分类结果 / 评测汇总 / SQL结果 / 来源 / 成本测算（公式实时计算）。
运行顺序：01 → 02 → 03 → 04 之后运行本脚本。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from common import ROOT, WINDOW_WEEKS, BRAND_WEEKLY, COST, load_ref, load_json, read_csv, review_state, quality_share

OUT = ROOT / 'docs' / '飞鹤投诉分拣案例_数据与测算.xlsx'
HEAD = PatternFill('solid', fgColor='1D1B17'); HEAD_FONT = Font(bold=True, color='F4EFE4')
INPUT = PatternFill('solid', fgColor='FFF2CC'); OUTF = PatternFill('solid', fgColor='D8E8E1')
THIN = Border(bottom=Side(style='thin', color='CFC5AD'))
WRAP = Alignment(wrap_text=True, vertical='top')


def table(ws, rows, widths=None, start=1, freeze=True):
    cols = list(rows[0])
    for j, c in enumerate(cols, 1):
        cell = ws.cell(start, j, c); cell.fill = HEAD; cell.font = HEAD_FONT; cell.alignment = WRAP
    for i, r in enumerate(rows, start + 1):
        for j, c in enumerate(cols, 1):
            v = r[c]
            cell = ws.cell(i, j, int(v) if isinstance(v, str) and v.isdigit() and len(v) < 6 else v)
            cell.alignment = WRAP; cell.border = THIN
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(start + 1, 1)
    return start + len(rows) + 1


ref = load_ref()
state = review_state(ref)
data = read_csv(ROOT / 'data' / 'complaints.csv')
n_q = sum(r['ref_category'] == '质量安全' for r in ref)
wb = Workbook()

# ---------- 说明 ----------
ws = wb.active; ws.title = '说明'
lines = [
    ('飞鹤业务研究与个人概念验证案例：公开投诉分拣（星飞帆系列）', True),
    ('范围：基于公开资料的飞鹤业务研究与个人概念验证，包括公开投诉的数据分析、关键词规则与 AI 的对比，以及单位经济测算；内部流程待访谈确认。', False),
    (f"数据：黑猫投诉 · 飞鹤商家页公开列表（未登录，2026-10-15 采集），筛选标题或摘要含「星飞帆」「卓睿」、发起日期 2026-07-15 至 2026-10-15，共 {len(data)} 条；按投诉单号去重；用户昵称未保存；{sum(r['truncated'] == '1' for r in data)} 条摘要被平台截断。", False),
    (f'口径：占比以这 {len(data)} 条公开投诉为分母，描述样本结构。', False),
    (f"标签（{state['short']}，{state['k']}/{state['n']}）：{state['long']} AI 结果由一个不读参考标签的 Claude 子代理离线批量生成，复核后不重新生成。", False),
    ('成本测算：黄色格为假设（时长、时薪、模型费、维护工时，可改），绿色格为按数据与标签计算的公式。', False),
    ('工作表：数据 / 分类结果 / 评测汇总 / SQL结果 / 来源 / 成本测算', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.font = Font(bold=b, size=13 if b else 11); c.alignment = WRAP
ws.column_dimensions['A'].width = 120

# ---------- 数据 ----------
ws = wb.create_sheet('数据')
keep = ['id', 'date', 'status', 'title', 'excerpt', 'demand', 'channel_mentioned', 'truncated', 'health_mention', 'batch_no', 'url']
table(ws, [{k: r[k] for k in keep} for r in data], [13, 11, 8, 28, 70, 18, 16, 8, 8, 15, 48])

# ---------- 分类结果 ----------
ws = wb.create_sheet('分类结果')
cases = read_csv(ROOT / 'outputs' / 'eval_cases.csv')
table(ws, cases, [13, 26, 12, 16, 12, 10, 10, 10, 10, 34, 8, 8, 18, 7, 7, 46])

# ---------- 评测汇总 ----------
ws = wb.create_sheet('评测汇总')
summ = load_json('outputs/eval_summary.json')
rows = []
for system in ['AI（Claude）', '关键词基线', '关键词基线v2（仅调换顺序）']:
    s = summ[system]
    rows.append({'方法': system, '类别与参考一致': s['类别一致'], '质量安全漏判': len(s['质量安全_漏判']),
                 '质量安全误判': '、'.join(s['质量安全_误判']) or 0, '健康症状漏判': len(s['健康症状_漏判']),
                 '需复核与参考一致': s['需复核_一致'], '标记需复核条数': s['标记需复核条数']})
end = table(ws, rows, [26, 16, 12, 30, 12, 16, 14])
ws.cell(end + 1, 1, f"AI 依据片段为原文逐字子串：{summ['AI 依据片段为原文逐字子串']}；基线只能给出命中的关键词，不给句子依据。")
ws.cell(end + 2, 1, f"复核状态：{summ['复核状态']}（{summ['已人工复核条数']}）。{summ['说明']}")

# ---------- SQL结果 ----------
ws = wb.create_sheet('SQL结果')
r0 = 1
for name, title in [('Q1_类别分布', f"Q1 样本中各类问题有多少（{state['short']}）"),
                    ('Q2_重复出现的具体问题', 'Q2 哪些具体问题在 ≥2 条独立投诉中重复出现'),
                    ('Q2b_活动权益里的规则类问题', 'Q2b 活动权益类里，规则/系统限制 vs 承诺未兑现'),
                    ('Q3_优先人工核查清单', 'Q3 哪些记录需要人工优先核查（排序分＝写到症状 3＋批次号 2＋多罐或多次 2＋未回复 1）')]:
    ws.cell(r0, 1, title).font = Font(bold=True, size=12)
    rows = read_csv(ROOT / 'outputs' / f'sql_{name}.csv')
    r0 = table(ws, rows, start=r0 + 1, freeze=False) + 2
for j, w in enumerate([14, 12, 14, 18, 10, 40, 48], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ---------- 来源 ----------
ws = wb.create_sheet('来源')
src = [
    ('飞鹤主要通过 2,800 余名线下客户（经销商）覆盖 7 万余个零售点；线下客户收入占乳制品收入 73.3%；产品也直接在大型电商平台及自有网站、App 销售', '官方披露', '中国飞鹤 2025 年度业绩公告（港交所）', 'https://www1.hkexnews.hk/listedco/listconews/sehk/2026/0326/2026032602610.pdf', '2026-03-26', 'p.20「Sales and Distribution Network」'),
    ('营销三要素：面对面研讨会（2025 年超 70 万场）、最大化线上互动、定向媒体曝光', '官方披露', '同上', '同上', '2026-03-26', 'p.21 营销策略'),
    ('年度业绩公告全文未出现 AI、artificial intelligence、digital、smart 等词', '作者检索（官方文件）', '同上（pdfplumber 全文检索）', '同上', '2026-03-26', '全文'),
    ('2026 上半年线下客户收入占乳制品收入 70.5%，约 33 万场面对面研讨会', '第三方整理的官方公告摘要', 'Reportify · 中国飞鹤 2026 中期业绩', 'https://reportify.cn/financials/1290110753197658112', '2026-08-26', '渠道段；港交所原文未取得'),
    ('2025 年营收 181.1 亿元、净利润 20.9 亿元；星妈会会员突破 8800 万；全龄营养转型', '第三方报道（企业信息）', '新华网', 'https://app.xinhuanet.com/news/article.html?articleId=202603274fd9f77a3d9d4b9e963d944d5d7e39d1', '2026-03-27', '正文'),
    ('依托火山引擎 HiAgent 搭建统一 AI 应用平台，已开发 436 个智能体；已应用场景为设备运维、膜清洗工艺、人事管理，未提消费者/客服/电商', '第三方报道（企业自述）', '鲁网（中国日报网转载）', 'https://cn.chinadaily.com.cn/a/202608/27/WS6a8ffbade4b09a165c7866ff.html', '2026-08-27', '正文'),
    ('内部 AI 平台接入 DeepSeek；「乳业 AI 场景地图」含电子商务、终端支持等，属规划，未说明上线情况', '第三方报道', '第一财经（新浪财经）', 'https://finance.sina.com.cn/roll/2025-02-20/doc-inemcpfe4627361.shtml', '2025-02-20', '正文'),
    ('客服智能质检覆盖 400 呼入、外呼与 IM 文本，全量质检；未提电商平台评价或公开投诉平台', '厂商案例（自述，未经第三方验证）', '沃丰科技官网', 'https://www.udesk.cn/ucm/report/52194', '2023-12-14', '正文'),
    ('2025 年投诉 48 小时结案率 94%；400 热线满意度 99%、星妈会在线 96%；风险框架新增「AI 技术应用风险」「直播业务风险」', '第三方转述官方 ESG 报告（原文未取得，待核）', '证券之星（网易号转载）', 'https://www.163.com/dy/article/KUG6HUGB051984TV.html', '2026-06-03', '客户服务段'),
    ('黑猫投诉飞鹤商家页显示近 30 天投诉量 52；本案例样本 41 条中列表显示「已回复」14 条', '公开数据（作者采集）', '黑猫投诉 · 飞鹤', 'https://tousu.sina.com.cn/company/view/?couid=1888345895', '2026-10-15 访问', '商家页头部与列表'),
    ('飞鹤内部如何处理公开投诉平台与电商评价、是否已用 AI 分拣', '尚未找到证据', '—', '—', '—', '列为访谈问题'),
]
table(ws, [dict(zip(['结论', '证据类型', '来源', '链接', '发布日期', '原文位置'], s)) for s in src], [60, 22, 24, 50, 12, 24])

# ---------- 成本测算 ----------
ws = wb.create_sheet('成本测算')
ws.column_dimensions['A'].width = 40; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 70
ws['A1'] = '净价值 ＝ 节省的人工处理成本 － 新增复核、模型及维护成本（按周）'; ws['A1'].font = Font(bold=True, size=12)
inputs = [
    ('每周新增条数', f"=(COUNTA(数据!A:A)-1)/{WINDOW_WEEKS}", f'实测（公式）：「数据」表样本条数 / {WINDOW_WEEKS} 周；飞鹤全部品牌约 {BRAND_WEEKLY} 条/周（160 条 / 14 周）'),
    ('人工：每条分拣＋摘依据＋填清单（分钟）', COST['manual_min'], '假设：读完一条约 1–2 分钟，判断与填表约 2 分钟；试点时实测替换'),
    ('工具辅助后：每条复核（分钟）', COST['review_min'], '假设：看类别与高亮依据，确认或改一下'),
    ('质量安全类占比（这类仍须人工通读，不计节省）', '=COUNTIF(分类结果!C:C,"质量安全")/(COUNTA(分类结果!A:A)-1)',
     f"实测（公式）：按「分类结果」表的最终参考类别计算，当前 {n_q}/{len(ref)}；{state['short']}"),
    ('人力成本（元/小时，含社保等）', COST['wage'], '假设'),
    ('模型费用（元/条）', COST['model_yuan'], '假设上限：每条约 600 输入＋150 输出 token，按国产通用模型公开价量级估'),
    ('维护与每月抽检（小时/月）', COST['maint_hours'], '假设：改提示词、规则与抽检 20 条'),
    ('一次性搭建（人日）', COST['build_days'], '假设：在现有智能体平台上配置，不含采购'),
]
ws['A3'], ws['B3'], ws['C3'] = '输入（黄色可改；绿色为按数据与标签计算的公式）', '数值', '依据'
for c in ('A3', 'B3', 'C3'):
    ws[c].fill = HEAD; ws[c].font = HEAD_FONT
for i, (k, v, note) in enumerate(inputs, 4):
    ws.cell(i, 1, k); ws.cell(i, 3, note).alignment = WRAP
    c = ws.cell(i, 2, v); c.fill = OUTF if isinstance(v, str) else INPUT
    if isinstance(v, str):
        c.number_format = '0.000'
# B4 条数 B5 人工分钟 B6 复核分钟 B7 质量占比 B8 时薪 B9 模型费 B10 维护小时 B11 搭建人日
outs = [
    ('每条可节省（元）', '=(1-B7)*(B5-B6)/60*B8', '只在非质量安全条目上节省'),
    ('节省的人工成本（元/周）', '=B4*B13', ''),
    ('新增模型费用（元/周）', '=B4*B9', ''),
    ('维护与抽检（元/周）', '=B10*B8/4.33', '每月按 4.33 周'),
    ('净价值（元/周）', '=B14-B15-B16', '负数＝不值得单独做'),
    ('盈亏平衡所需条数（条/周）', '=B16/(B13-B9)', '每周新增条数超过它，净价值才为正'),
    ('一次性搭建成本（元）', '=B11*8*B8', ''),
    ('搭建成本回收（周）', '=IF(B17>0,B19/B17,"净价值≤0，不回收")', ''),
]
ws['A12'], ws['B12'] = '结果（绿色＝公式）', ''
ws['A12'].fill = HEAD; ws['A12'].font = HEAD_FONT; ws['B12'].fill = HEAD
for i, (k, f, note) in enumerate(outs, 13):
    ws.cell(i, 1, k); c = ws.cell(i, 2, f); c.fill = OUTF; c.number_format = '0.00'; ws.cell(i, 3, note)
ws['A23'] = '敏感性：不同周量下的净价值（其他假设同上）'; ws['A23'].font = Font(bold=True)
ws['A24'], ws['B24'], ws['C24'] = '每周新增条数', '净价值（元/周）', '对应情形'
for c in ('A24', 'B24', 'C24'):
    ws[c].fill = HEAD; ws[c].font = HEAD_FONT
scen = [('=B4', '本系列公开投诉（实测，随数据变化）'), (BRAND_WEEKLY, '飞鹤全部品牌公开投诉（实测）'), (50, '假设：再加一个电商平台的差评'),
        (200, '假设：多平台差评＋客服工单'), (1000, '假设：全渠道评价与工单')]
for i, (n, s) in enumerate(scen, 25):
    ws.cell(i, 1, n).number_format = '0.0'; c = ws.cell(i, 2, f'=A{i}*($B$13-$B$9)-$B$16'); c.fill = OUTF; c.number_format = '0.00'; ws.cell(i, 3, s)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print('写出：', OUT)
