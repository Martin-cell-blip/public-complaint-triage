# -*- coding: utf-8 -*-
"""参考标签人工复核。
  python src/07_review_sheet.py make [--force]   生成 labels/参考标签复核表.xlsx；文件已存在时拒绝覆盖（防止冲掉已填内容），确需重建加 --force
  python src/07_review_sheet.py apply            只回写「本行已复核＝是」的行，这些行标为「已人工复核」；其余行原样保留
  python src/07_review_sheet.py apply --confirm-all
        整表视为已逐条复核：空白＝同意当前标签，全部标为「已人工复核」。
        只有在复核人明确确认「全表已逐条看完」之后才能使用这个参数。
复核表只放原文与当前参考标签，不放 AI 与规则结果，避免复核时被带偏。
本脚本只改 labels/reference_labels.csv；AI 与关键词规则的预测文件不动，也不会因人工答案重新生成。
"""
import csv, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from common import ROOT, load_ref, read_csv

REF = ROOT / 'labels' / 'reference_labels.csv'
SHEET = ROOT / 'labels' / '参考标签复核表.xlsx'
CATS = ['质量安全', '活动权益', '订单发货', '其他', '信息不足']
COLS = ['id', '投诉原文（标题。摘要）', '当前类别', '当前_写到症状', '当前_需复核', '当前子问题', '当前备注',
        '复核_类别', '复核_写到症状', '复核_需复核', '复核意见', '本行已复核']
EDIT = ['复核_类别', '复核_写到症状', '复核_需复核', '复核意见']


def make(force):
    if SHEET.exists() and not force:
        sys.exit(f'{SHEET} 已存在，未覆盖。确认里面没有要保留的复核内容后，再加 --force 重建。')
    texts = {r['id']: r for r in read_csv(ROOT / 'data' / 'complaints.csv')}
    ref = load_ref()
    wb = Workbook(); ws = wb.active; ws.title = '复核'
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color='F4EFE4'); c.fill = PatternFill('solid', fgColor='1D1B17'); c.alignment = Alignment(wrap_text=True)
    for r in ref:
        ws.append([r['id'], texts[r['id']]['text'], r['ref_category'], '是' if r['ref_health'] == '1' else '否',
                   '是' if r['ref_needs_review'] == '1' else '否', r['subtag'], r['note'], '', '', '', '',
                   '是' if r['label_status'] == '已人工复核' else '否'])
    n = len(ref) + 1
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        for c in row[7:]:
            c.fill = PatternFill('solid', fgColor='FFF2CC')
    dv_cat = DataValidation(type='list', formula1='"' + ','.join(CATS) + '"', allow_blank=True)
    dv_yn = DataValidation(type='list', formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv_cat); ws.add_data_validation(dv_yn)
    dv_cat.add(f'H2:H{n}'); dv_yn.add(f'I2:J{n}'); dv_yn.add(f'L2:L{n}')
    for col, w in zip('ABCDEFGHIJKL', [13, 70, 10, 8, 8, 14, 30, 11, 9, 9, 30, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C2'
    g = wb.create_sheet('类别定义')
    for line in [
        '质量安全：奶粉本身的问题——异物、虫、异味、结块、变质或临期、罐内配件错误（如量勺错配）、食用后身体不适',
        '活动权益：集罐、买赠、育儿补贴、扫码积分、赠品（含赠品本身的质量）、权益卡、优惠券等活动规则或兑现',
        '订单发货：下单后拒发、订单异常、拦截、取消订单、限购、超时未发货',
        '其他：与上面三类都无关；信息不足：文字太短或被截断，看不出核心问题',
        '多个问题并存时，选投诉人主要诉求针对的问题；写到食用后症状的，「写到症状」一律为是。',
        '需复核：质量安全一律是；写到症状一律是；类别拿不准或多个问题并存时是；其余否。',
        '复核方法：黄色的复核列留空＝同意当前标签；不同意就选下拉值，并在「复核意见」写一句理由。',
        '「本行已复核」：看完一行就改成「是」，可以分批进行。只有标了「是」的行才会被回写并计入「已人工复核」。',
    ]:
        g.append([line])
    g.column_dimensions['A'].width = 110
    wb.save(SHEET)
    print('写出：', SHEET)


def apply(confirm_all):
    ws = load_workbook(SHEET)['复核']
    rows = {str(r[0]): dict(zip(COLS, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    ref = load_ref()
    if set(rows) != {r['id'] for r in ref}:
        sys.exit('复核表与 reference_labels.csv 的投诉单号不一致，未写回。')
    done = {i for i, v in rows.items() if confirm_all or v['本行已复核'] == '是'}
    pending_edits = sorted(i for i, v in rows.items() if i not in done and any(v[c] for c in EDIT))
    if pending_edits:
        sys.exit('这些行填了复核内容，但「本行已复核」不是「是」，无法判断是否已复核完，未写回：' + '、'.join(pending_edits))

    yn = {'是': '1', '否': '0'}
    changes = []
    for r in ref:
        if r['id'] not in done:
            continue
        v = rows[r['id']]
        new = {'ref_category': v['复核_类别'] or r['ref_category'],
               'ref_health': yn[v['复核_写到症状']] if v['复核_写到症状'] else r['ref_health'],
               'ref_needs_review': yn[v['复核_需复核']] if v['复核_需复核'] else r['ref_needs_review']}
        diff = [f'{k}: {r[k]}→{new[k]}' for k in new if new[k] != r[k]]
        if diff:
            changes.append(f"{r['id']} {'；'.join(diff)} ｜ {v['复核意见'] or ''}")
        r.update(new)
        if v['复核意见'] and f"（复核：{v['复核意见']}）" not in r['note']:
            r['note'] = f"{r['note']}（复核：{v['复核意见']}）"
        r['label_status'] = '已人工复核'
    with REF.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(ref[0])); w.writeheader(); w.writerows(ref)
    print('\n'.join(changes) if changes else '没有改动类别、症状或需复核标记')
    k = sum(r['label_status'] == '已人工复核' for r in ref)
    print(f'已写回 {REF}：本次回写 {len(done)} 行，改动 {len(changes)} 行；累计已人工复核 {k}/{len(ref)}')


cmd, flags = sys.argv[1], set(sys.argv[2:])
if cmd == 'make':
    make('--force' in flags)
elif cmd == 'apply':
    apply('--confirm-all' in flags)
else:
    sys.exit('用法：make [--force] | apply [--confirm-all]')
